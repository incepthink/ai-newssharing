import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Set padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    """Set custom borders on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders_elm)

def create_security_document():
    doc = docx.Document()

    # Color Palette
    NAVY = RGBColor(15, 41, 66)        # #0F2942 - Primary Brand Navy
    ACCENT_BLUE = RGBColor(29, 99, 237) # #1D63ED - Vibrant Blue Accent
    DARK_SLATE = RGBColor(51, 65, 85)   # #334155 - Body Text Charcoal
    MUTED_GRAY = RGBColor(100, 116, 139)# #64748B - Muted Subtitles & Meta
    WHITE = RGBColor(255, 255, 255)
    
    # Page setup - Margins (0.75 in all sides for clean executive density)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # HEADER / MASTHEAD TABLE
    # ----------------------------------------------------
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    col_widths = [Inches(4.8), Inches(2.2)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Left cell: Brand & Document Title
    left_cell = header_table.cell(0, 0)
    p_brand = left_cell.paragraphs[0]
    p_brand.paragraph_format.space_after = Pt(2)
    r_brand = p_brand.add_run("HASHCASE  |  CLIENT DATA PRIVACY & SECURITY BRIEF")
    r_brand.font.size = Pt(8.5)
    r_brand.font.bold = True
    r_brand.font.color.rgb = ACCENT_BLUE

    p_title = left_cell.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("How HashCase Protects Your Data & IP")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = left_cell.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("Practical, Hardened Security for Your Custom AI Workflows")
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = MUTED_GRAY

    # Right cell: Metadata Badge
    right_cell = header_table.cell(0, 1)
    set_cell_background(right_cell, "F8FAFC")
    set_cell_borders(right_cell, top="single", bottom="single", left="single", right="single", color="E2E8F0", sz="4")
    set_cell_margins(right_cell, top=70, bottom=70, left=100, right=100)
    
    p_meta = right_cell.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(1)
    r_m1 = p_meta.add_run("TARGET: ")
    r_m1.font.bold = True
    r_m1.font.size = Pt(7.5)
    r_m1.font.color.rgb = MUTED_GRAY
    r_m2 = p_meta.add_run("Founders, CTOs & Ops Leaders\n")
    r_m2.font.size = Pt(7.5)
    
    r_m3 = p_meta.add_run("SCOPE: ")
    r_m3.font.bold = True
    r_m3.font.size = Pt(7.5)
    r_m3.font.color.rgb = MUTED_GRAY
    r_m4 = p_meta.add_run("Data Storage, Privacy & Fine-Tuning\n")
    r_m4.font.size = Pt(7.5)

    r_m5 = p_meta.add_run("STANDARD: ")
    r_m5.font.bold = True
    r_m5.font.size = Pt(7.5)
    r_m5.font.color.rgb = MUTED_GRAY
    r_m6 = p_meta.add_run("Zero-Leakage Private Model Hosting")
    r_m6.font.bold = True
    r_m6.font.size = Pt(7.5)
    r_m6.font.color.rgb = ACCENT_BLUE

    # Divider bar
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    div_table = doc.add_table(rows=1, cols=1)
    div_cell = div_table.cell(0, 0)
    div_cell.width = Inches(7.0)
    set_cell_background(div_cell, "1D63ED")
    set_cell_margins(div_cell, top=10, bottom=10, left=0, right=0)
    p_div = div_cell.paragraphs[0]
    p_div.paragraph_format.space_after = Pt(0)
    r_div = p_div.add_run("")
    r_div.font.size = Pt(1)

    # ----------------------------------------------------
    # EXECUTIVE VALUE STATEMENT (Callout Box)
    # ----------------------------------------------------
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout.autofit = False
    c_cell = callout.cell(0, 0)
    c_cell.width = Inches(7.0)
    set_cell_background(c_cell, "F1F5F9")
    set_cell_borders(c_cell, left="single", color="1D63ED", sz="16")
    set_cell_margins(c_cell, top=80, bottom=80, left=140, right=120)

    p_callout = c_cell.paragraphs[0]
    p_callout.paragraph_format.space_after = Pt(2)
    p_callout.paragraph_format.line_spacing = 1.15
    r_c_bold = p_callout.add_run("The HashCase Security Commitment: ")
    r_c_bold.font.bold = True
    r_c_bold.font.color.rgb = NAVY
    r_c_bold.font.size = Pt(9.5)
    r_c_text = p_callout.add_run(
        "When you build custom AI automations with HashCase, your business data is never shared, never leaked, "
        "and never used to benefit another company. We take leading open-source models and fine-tune them exclusively "
        "for your business. Your model weights and proprietary records are hosted in dedicated, isolated environments "
        "with closed network ports and banking-grade encryption from day one."
    )
    r_c_text.font.size = Pt(9.0)
    r_c_text.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # THE 4 CORE SECURITY PILLARS (Practical & On-The-Ground)
    # ----------------------------------------------------
    p_sec_title = doc.add_paragraph()
    p_sec_title.paragraph_format.space_before = Pt(6)
    p_sec_title.paragraph_format.space_after = Pt(3)
    r_s = p_sec_title.add_run("Our 4 Practical Security & Data Privacy Pillars")
    r_s.font.size = Pt(11.5)
    r_s.font.bold = True
    r_s.font.color.rgb = NAVY

    pillars = [
        {
            "num": "1",
            "title": "Private Open-Source Models: No Leaks, Total Ownership",
            "subtitle": "Your Data Fine-Tunes Models That ONLY Your Company Uses",
            "what_we_do": (
                "Unlike typical agencies that send your raw customer and business records to public black-box APIs, "
                "we fine-tune high-performance open-source models (such as Meta's Llama 3, Mistral, or Qwen) "
                "specifically on your company's workflows. These fine-tuned model weights are deployed in dedicated, private "
                "containers. Your data is NEVER pooled, NEVER shared, and NEVER used to train models for anyone else. "
                "Your competitors will never benefit from your data, and your custom model intelligence remains 100% yours."
            ),
            "guarantee": "Zero data leakage, absolute model exclusivity, and 100% intellectual property ownership."
        },
        {
            "num": "2",
            "title": "Hardened Storage Architecture: Unbreachable Cloud Infrastructure",
            "subtitle": "Zero Public Exposure, Dedicated Databases & Banking-Grade Encryption",
            "what_we_do": (
                "We do not store your information in open buckets or expose databases to the internet. We deploy your "
                "storage inside an isolated Virtual Private Cloud (VPC) with all public database ports strictly closed—accessible "
                "only by authenticated, authorized application code. All data is protected with AES-256 encryption at rest "
                "(rendering disk data unreadable gibberish if physically compromised) and TLS/HTTPS encryption in transit. "
                "Furthermore, we enforce strict multi-tenant isolation: your data is never mixed with another client's records."
            ),
            "guarantee": "Closed-port network isolation, AES-256 encryption, and zero cross-company data co-mingling."
        },
        {
            "num": "3",
            "title": "Ground-Level Operational Hygiene: Real Security in Practice",
            "subtitle": "Strict Key Vaults, Least-Privilege Access & Prompt Injection Defenses",
            "what_we_do": (
                "Real security lives in everyday engineering discipline. We store all database credentials, API keys, and tokens "
                "in encrypted cloud secret managers (like AWS Secrets Manager)—never hardcoded in code. We enforce the Principle of "
                "Least Privilege: only the specific engineers directly building your pipeline have temporary access, protected by "
                "mandatory Multi-Factor Authentication (MFA). We also implement prompt injection and extraction firewalls so outside users "
                "cannot trick your AI into revealing internal system prompts, database keys, or private customer records."
            ),
            "guarantee": "No unauthorized staff browsing, zero exposed credentials, and hardened defenses against AI exploitation."
        },
        {
            "num": "4",
            "title": "Business Continuity & Your Exit Freedom",
            "subtitle": "Automated Daily Backups with a 100% Permanent Deletion Guarantee",
            "what_we_do": (
                "We protect your operations with automated daily encrypted backups of your databases and model checkpoints, "
                "ensuring a server crash or accidental deletion can never halt your business. Equally important is our exit policy: "
                "we never hold your data or models hostage. If our engagement ever concludes, all raw documents, fine-tuned weights, "
                "vector embeddings, and database snapshots are permanently wiped from our systems upon your request, with zero residual copies."
            ),
            "guarantee": "Continuous operational uptime with zero vendor lock-in and a verifiable, complete data deletion guarantee."
        }
    ]

    for p in pillars:
        t_p = doc.add_table(rows=1, cols=1)
        t_p.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_p.autofit = False
        cell_p = t_p.cell(0, 0)
        cell_p.width = Inches(7.0)
        set_cell_background(cell_p, "FFFFFF")
        set_cell_borders(cell_p, left="single", color="CBD5E1", sz="8")
        set_cell_margins(cell_p, top=35, bottom=35, left=90, right=70)

        p_h = cell_p.paragraphs[0]
        p_h.paragraph_format.space_before = Pt(0)
        p_h.paragraph_format.space_after = Pt(1)
        
        r_num = p_h.add_run(f"PILLAR {p['num']}: ")
        r_num.font.bold = True
        r_num.font.size = Pt(8.5)
        r_num.font.color.rgb = ACCENT_BLUE

        r_title = p_h.add_run(p['title'])
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = NAVY

        p_sub = cell_p.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(2)
        r_sub = p_sub.add_run(p['subtitle'])
        r_sub.font.italic = True
        r_sub.font.size = Pt(8.0)
        r_sub.font.color.rgb = MUTED_GRAY

        p_desc = cell_p.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(2)
        p_desc.paragraph_format.line_spacing = 1.12
        r_do_lbl = p_desc.add_run("What We Do: ")
        r_do_lbl.font.bold = True
        r_do_lbl.font.size = Pt(8.5)
        r_do_lbl.font.color.rgb = DARK_SLATE
        r_do_txt = p_desc.add_run(p['what_we_do'])
        r_do_txt.font.size = Pt(8.5)
        r_do_txt.font.color.rgb = DARK_SLATE

        p_guar = cell_p.add_paragraph()
        p_guar.paragraph_format.space_before = Pt(0)
        p_guar.paragraph_format.space_after = Pt(0)
        r_g_lbl = p_guar.add_run("Client Guarantee: ")
        r_g_lbl.font.bold = True
        r_g_lbl.font.size = Pt(8.0)
        r_g_lbl.font.color.rgb = ACCENT_BLUE
        r_g_txt = p_guar.add_run(p['guarantee'])
        r_g_txt.font.italic = True
        r_g_txt.font.size = Pt(8.0)
        r_g_txt.font.color.rgb = NAVY

        # Subtle spacing between pillars
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(1)

    # ----------------------------------------------------
    # COMPARISON TABLE: HashCase vs Typical AI Agency
    # ----------------------------------------------------
    p_mat_title = doc.add_paragraph()
    p_mat_title.paragraph_format.space_before = Pt(4)
    p_mat_title.paragraph_format.space_after = Pt(3)
    r_mat = p_mat_title.add_run("How HashCase Compares to Typical AI Agencies")
    r_mat.font.size = Pt(11.0)
    r_mat.font.bold = True
    r_mat.font.color.rgb = NAVY

    matrix_data = [
        ("Security Factor", "Typical AI Agency / Off-The-Shelf Tools", "The HashCase Standard"),
        ("Model Hosting", "Data sent to public third-party APIs (black-box)", "Private open-source models hosted in dedicated containers"),
        ("Model Ownership", "You rent access; vendor owns the intelligence", "You own the fine-tuned model and custom weights"),
        ("Data Storage", "Shared multi-tenant databases with open ports", "Isolated VPC subnets, closed ports & strict tenant separation"),
        ("Encryption", "Basic or default settings, often incomplete", "AES-256 at rest and TLS/HTTPS in transit everywhere"),
        ("Competitor Risk", "Risk of your data training general vendor models", "Zero-leakage guarantee; your data is never pooled or shared"),
        ("Access & Keys", "Keys often hardcoded in scripts or shared drives", "Encrypted secret vaults (AWS Secrets Manager) & MFA"),
        ("Exit & Deletion", "Proprietary lock-in with residual data retention", "Complete data export or permanent cryptographic wipe")
    ]

    mat_table = doc.add_table(rows=len(matrix_data), cols=3)
    mat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mat_table.autofit = False

    t_widths = [Inches(1.6), Inches(2.7), Inches(2.7)]

    for r_idx, row_data in enumerate(matrix_data):
        row = mat_table.rows[r_idx]
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = t_widths[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.space_after = Pt(0)
            p_cell.paragraph_format.line_spacing = 1.05

            if r_idx == 0:
                # Header row
                set_cell_background(cell, "0F2942")
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="0F2942", sz="6")
                r = p_cell.add_run(text)
                r.font.bold = True
                r.font.size = Pt(8.0)
                r.font.color.rgb = WHITE
            else:
                # Data rows
                bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="E2E8F0", sz="4")
                r = p_cell.add_run(text)
                r.font.size = Pt(7.5)
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = NAVY
                elif c_idx == 2:
                    r.font.bold = True
                    r.font.color.rgb = NAVY
                else:
                    r.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # BOTTOM CALLOUT / CALL TO ACTION
    # ----------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    cta_table = doc.add_table(rows=1, cols=1)
    cta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cta_table.autofit = False
    cta_cell = cta_table.cell(0, 0)
    cta_cell.width = Inches(7.0)
    set_cell_background(cta_cell, "0F2942")
    set_cell_margins(cta_cell, top=70, bottom=70, left=120, right=120)

    p_cta1 = cta_cell.paragraphs[0]
    p_cta1.paragraph_format.space_after = Pt(1)
    r_cta1 = p_cta1.add_run("Built for Real Businesses. Protected by Real Engineering.")
    r_cta1.font.bold = True
    r_cta1.font.size = Pt(9.5)
    r_cta1.font.color.rgb = WHITE

    p_cta2 = cta_cell.add_paragraph()
    p_cta2.paragraph_format.space_after = Pt(2)
    p_cta2.paragraph_format.line_spacing = 1.1
    r_cta2 = p_cta2.add_run(
        "You don't need an unreadable 50-page corporate manual to know your data is safe. With HashCase, "
        "your data stays in your private environment, tunes your private models, and powers your business—with "
        "zero leaks, zero public sharing, and 100% data ownership."
    )
    r_cta2.font.size = Pt(8.0)
    r_cta2.font.color.rgb = RGBColor(226, 232, 240)

    p_cta3 = cta_cell.add_paragraph()
    p_cta3.paragraph_format.space_after = Pt(0)
    r_cta3_bold = p_cta3.add_run("Have questions about our security setup? ")
    r_cta3_bold.font.bold = True
    r_cta3_bold.font.size = Pt(8.0)
    r_cta3_bold.font.color.rgb = RGBColor(147, 197, 253)
    r_cta3_text = p_cta3.add_run("Contact HashCase Engineering directly: ")
    r_cta3_text.font.size = Pt(8.0)
    r_cta3_text.font.color.rgb = WHITE
    r_contact = p_cta3.add_run("hello@hashcase.com  |  www.hashcase.com")
    r_contact.font.bold = True
    r_contact.font.size = Pt(8.0)
    r_contact.font.color.rgb = RGBColor(147, 197, 253)

    # Save document
    output_path = r"c:\Users\shaik\Desktop\dev-work\ai-newssharing\HashCase_Client_Data_Security_Brief.docx"
    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    create_security_document()
