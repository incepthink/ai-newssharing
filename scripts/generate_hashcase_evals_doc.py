import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if present
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    """Set custom borders on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders_elm)

def create_document():
    doc = docx.Document()

    # Define color palette
    NAVY = RGBColor(15, 41, 66)        # #0F2942 - Primary Brand
    ACCENT_BLUE = RGBColor(29, 99, 237) # #1D63ED - Vibrant Accent
    DARK_SLATE = RGBColor(51, 65, 85)   # #334155 - Body Text
    MUTED_GRAY = RGBColor(100, 116, 139)# #64748B - Subtitles & Meta
    WHITE = RGBColor(255, 255, 255)
    
    # Page setup - Margins (0.75 in all sides for clean executive density)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base Style Configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # HEADER / MASTHEAD TABLE (Clean branded layout)
    # ----------------------------------------------------
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    col_widths = [Inches(4.8), Inches(2.2)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Left cell: Brand & Document Title
    left_cell = header_table.cell(0, 0)
    p_brand = left_cell.paragraphs[0]
    p_brand.paragraph_format.space_after = Pt(2)
    r_brand = p_brand.add_run("HASHCASE  |  ENTERPRISE AI CAPABILITY BRIEF")
    r_brand.font.size = Pt(8.5)
    r_brand.font.bold = True
    r_brand.font.color.rgb = ACCENT_BLUE

    p_title = left_cell.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Enterprise AI Workflow Evaluation Framework")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = left_cell.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run("How We Benchmark, Validate, and Guarantee Production-Grade Automations")
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = MUTED_GRAY

    # Right cell: Metadata Badge
    right_cell = header_table.cell(0, 1)
    set_cell_background(right_cell, "F8FAFC")
    set_cell_borders(right_cell, top="single", bottom="single", left="single", right="single", color="E2E8F0", sz="4")
    set_cell_margins(right_cell, top=80, bottom=80, left=120, right=120)
    
    p_meta = right_cell.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(1)
    r_m1 = p_meta.add_run("TARGET: ")
    r_m1.font.bold = True
    r_m1.font.size = Pt(7.5)
    r_m1.font.color.rgb = MUTED_GRAY
    r_m2 = p_meta.add_run("Enterprise Operations & Tech Leadership\n")
    r_m2.font.size = Pt(7.5)
    
    r_m3 = p_meta.add_run("SCOPE: ")
    r_m3.font.bold = True
    r_m3.font.size = Pt(7.5)
    r_m3.font.color.rgb = MUTED_GRAY
    r_m4 = p_meta.add_run("Automated Workflow Reliability (Evals)\n")
    r_m4.font.size = Pt(7.5)

    r_m5 = p_meta.add_run("STATUS: ")
    r_m5.font.bold = True
    r_m5.font.size = Pt(7.5)
    r_m5.font.color.rgb = MUTED_GRAY
    r_m6 = p_meta.add_run("Production Hardening Standard")
    r_m6.font.bold = True
    r_m6.font.size = Pt(7.5)
    r_m6.font.color.rgb = ACCENT_BLUE

    # Divider bar
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    div_table = doc.add_table(rows=1, cols=1)
    div_cell = div_table.cell(0, 0)
    div_cell.width = Inches(7.0)
    set_cell_background(div_cell, "1D63ED")
    set_cell_margins(div_cell, top=10, bottom=10, left=0, right=0)
    p_div = div_cell.paragraphs[0]
    p_div.paragraph_format.space_after = Pt(0)
    r_div = p_div.add_run("")
    r_div.font.size = Pt(1)

    # ----------------------------------------------------
    # EXECUTIVE VALUE STATEMENT (Callout Box)
    # ----------------------------------------------------
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout.autofit = False
    c_cell = callout.cell(0, 0)
    c_cell.width = Inches(7.0)
    set_cell_background(c_cell, "F1F5F9")
    set_cell_borders(c_cell, left="single", color="1D63ED", sz="16") # Thick blue accent left border
    set_cell_margins(c_cell, top=100, bottom=100, left=160, right=140)

    p_callout = c_cell.paragraphs[0]
    p_callout.paragraph_format.space_after = Pt(3)
    p_callout.paragraph_format.line_spacing = 1.15
    r_c_bold = p_callout.add_run("The HashCase Guarantee: ")
    r_c_bold.font.bold = True
    r_c_bold.font.color.rgb = NAVY
    r_c_bold.font.size = Pt(9.5)
    r_c_text = p_callout.add_run(
        "General-purpose LLMs are inherently probabilistic, but enterprise operations require 100% deterministic guarantees. "
        "At HashCase, we do not deploy raw prompts into your organization. We engineer, stress-test, and deploy "
        "hardened AI workflows governed by a 6-pillar automated evaluation framework. Every automated action is validated "
        "for business logic correctness, API precision, cost predictability, and security before touching client systems."
    )
    r_c_text.font.size = Pt(9.0)
    r_c_text.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # THE 6 CORE EVALUATION PILLARS
    # ----------------------------------------------------
    p_sec_title = doc.add_paragraph()
    p_sec_title.paragraph_format.space_before = Pt(8)
    p_sec_title.paragraph_format.space_after = Pt(4)
    r_s = p_sec_title.add_run("Our 6 Core Evaluation Pillars")
    r_s.font.size = Pt(12)
    r_s.font.bold = True
    r_s.font.color.rgb = NAVY

    pillars = [
        {
            "num": "1",
            "title": "Deterministic & Programmatic Business Logic Evals",
            "subtitle": "Eliminating Silent Failure through Code-Level Assertions",
            "what_we_do": (
                "We do not rely on AI to verify its own factual or arithmetic correctness. Before any workflow output "
                "or database update is approved, our automated harness runs strict code-level assertions: validating "
                "strict data schemas (JSON/Pydantic), confirming relational invariants (e.g., invoice totals = subtotal + tax), "
                "and verifying foreign keys and entity IDs directly against your enterprise records."
            ),
            "guarantee": "Zero hallucinated schema fields, 100% mathematical consistency, and guaranteed type safety."
        },
        {
            "num": "2",
            "title": "Multi-Step & Agent Trajectory Evals",
            "subtitle": "Tool Selection Accuracy, Trajectory Efficiency & Loop Detection",
            "what_we_do": (
                "Modern automations chain multiple tools (CRMs, ERPs, APIs, databases) together. We evaluate the entire "
                "execution path: measuring tool selection precision/recall (ensuring the agent selects the exact right tool with "
                "correct arguments), tracking trajectory efficiency (eliminating wasteful intermediate calls), and enforcing "
                "deterministic loop detection to ensure an agent never gets trapped in infinite execution cycles."
            ),
            "guarantee": "Predictable, streamlined agent execution without runaway API loops, wasted tokens, or misfired webhooks."
        },
        {
            "num": "3",
            "title": "LLM-as-a-Judge with Calibrated Rubrics",
            "subtitle": "Automated Qualitative Audits Calibrated to Your Senior Leadership",
            "what_we_do": (
                "For qualitative outputs (customer emails, policy summaries, case triage), we deploy frontier LLM judges "
                "governed by exhaustive, multi-point rubrics (evaluating faithfulness to source documents, tone, completeness, "
                "and policy compliance). Crucially, we calibrate our automated judges against your human Subject Matter Experts (SMEs) "
                "to ensure our automated evaluations achieve >90% statistical alignment with your senior staff."
            ),
            "guarantee": "Automated, scalable quality control that grades every output with the exact standards of your best employees."
        },
        {
            "num": "4",
            "title": "Sandbox Simulation & Mock Environments",
            "subtitle": "Safe, Side-Effect-Free Validation Against Enterprise Systems",
            "what_we_do": (
                "Testing live AI agents against real production databases or ticketing platforms creates immense operational risk. "
                "HashCase evaluates workflows in isolated sandboxes and mock API environments. We validate state transitions—confirming "
                "not just what the model outputs, but that the mock CRM record, inventory count, or status flag transitioned to the "
                "exact intended state under simulated network delays, rate limits, and service outages."
            ),
            "guarantee": "Zero risk to production data, with proven state-level verification before live deployment."
        },
        {
            "num": "5",
            "title": "Curated 'Golden Datasets' & Synthetic Expansion",
            "subtitle": "Continuous Stress-Testing Against Adversarial Edge Cases",
            "what_we_do": (
                "We curate a proprietary regression test suite from your historical enterprise records and real-world edge cases. "
                "We then artificially expand this dataset with synthetic adversarial generation: intentionally injecting typos, "
                "missing fields, contradictory customer requests, and malformed inputs. Crucially, every production anomaly or corner-case "
                "discovered is converted into a permanent regression test case."
            ),
            "guarantee": "An automation system that becomes systematically more resilient with every iteration, preventing regression."
        },
        {
            "num": "6",
            "title": "Operational, Cost & Security Evals",
            "subtitle": "Latency Budgets, Token Attribution & Adversarial Jailbreak Defense",
            "what_we_do": (
                "Enterprise readiness requires meeting strict operational SLAs. We benchmark Time-to-First-Token (TTFT) and "
                "P95/P99 execution latency. We enforce granular cost-per-workflow attribution to ensure high ROI and intelligent "
                "model routing. Furthermore, we execute automated security red-teaming to defend against indirect prompt injection, "
                "jailbreaks, and unauthorized PII leakage."
            ),
            "guarantee": "Predictable cloud unit economics, guaranteed SLA compliance, and enterprise-grade data privacy."
        }
    ]

    for p in pillars:
        # Pillar container table for clean indentation and subtle border
        t_p = doc.add_table(rows=1, cols=1)
        t_p.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_p.autofit = False
        cell_p = t_p.cell(0, 0)
        cell_p.width = Inches(7.0)
        set_cell_background(cell_p, "FFFFFF")
        set_cell_borders(cell_p, left="single", color="CBD5E1", sz="8")
        set_cell_margins(cell_p, top=40, bottom=40, left=100, right=80)

        p_h = cell_p.paragraphs[0]
        p_h.paragraph_format.space_before = Pt(0)
        p_h.paragraph_format.space_after = Pt(2)
        
        # Pillar number badge & title
        r_num = p_h.add_run(f"PILLAR {p['num']}: ")
        r_num.font.bold = True
        r_num.font.size = Pt(8.5)
        r_num.font.color.rgb = ACCENT_BLUE

        r_title = p_h.add_run(p['title'])
        r_title.font.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = NAVY

        # Subtitle
        p_sub = cell_p.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(2)
        r_sub = p_sub.add_run(p['subtitle'])
        r_sub.font.italic = True
        r_sub.font.size = Pt(8.5)
        r_sub.font.color.rgb = MUTED_GRAY

        # What We Do
        p_desc = cell_p.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(2)
        p_desc.paragraph_format.line_spacing = 1.15
        r_do_lbl = p_desc.add_run("What We Do: ")
        r_do_lbl.font.bold = True
        r_do_lbl.font.size = Pt(9.0)
        r_do_lbl.font.color.rgb = DARK_SLATE
        r_do_txt = p_desc.add_run(p['what_we_do'])
        r_do_txt.font.size = Pt(9.0)
        r_do_txt.font.color.rgb = DARK_SLATE

        # Client Guarantee
        p_guar = cell_p.add_paragraph()
        p_guar.paragraph_format.space_before = Pt(1)
        p_guar.paragraph_format.space_after = Pt(1)
        r_g_lbl = p_guar.add_run("Client Guarantee: ")
        r_g_lbl.font.bold = True
        r_g_lbl.font.size = Pt(8.5)
        r_g_lbl.font.color.rgb = ACCENT_BLUE
        r_g_txt = p_guar.add_run(p['guarantee'])
        r_g_txt.font.italic = True
        r_g_txt.font.size = Pt(8.5)
        r_g_txt.font.color.rgb = NAVY

        # Spacing between pillars
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(1)

    # ----------------------------------------------------
    # SUMMARY COMPARISON TABLE
    # ----------------------------------------------------
    p_mat_title = doc.add_paragraph()
    p_mat_title.paragraph_format.space_before = Pt(6)
    p_mat_title.paragraph_format.space_after = Pt(4)
    r_mat = p_mat_title.add_run("HashCase Evaluation Architecture at a Glance")
    r_mat.font.size = Pt(11.5)
    r_mat.font.bold = True
    r_mat.font.color.rgb = NAVY

    matrix_data = [
        ("Evaluation Layer", "Primary Metrics & Checks", "HashCase Safeguard for Your Business"),
        ("Deterministic & Logic", "Schema validation, arithmetic consistency, foreign key matching", "Zero corrupted database records or malformed payloads"),
        ("Agent Trajectory", "Tool selection accuracy, step-count efficiency, loop breaker", "Eliminates runaway API costs and infinite retry cycles"),
        ("LLM-as-a-Judge", "Groundedness/faithfulness, tone, policy compliance, SME rubrics", "Continuous auditing matching senior leadership's standards"),
        ("Sandbox Simulation", "Mock CRM/ERP API calls, state transitions, failure recovery", "Guarantees safe execution with zero production risk"),
        ("Golden Datasets", "SME historical baselines, synthetic adversarial stress cases", "Prevents regressions; automations strengthen over time"),
        ("Operational & Security", "P95/P99 latency, TTFT, cost-per-run, prompt injection defense", "Ensures strict SLA compliance, budget control & data privacy")
    ]

    mat_table = doc.add_table(rows=len(matrix_data), cols=3)
    mat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mat_table.autofit = False

    t_widths = [Inches(1.8), Inches(2.7), Inches(2.5)]

    for r_idx, row_data in enumerate(matrix_data):
        row = mat_table.rows[r_idx]
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = t_widths[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.space_after = Pt(0)
            p_cell.paragraph_format.line_spacing = 1.1

            if r_idx == 0:
                # Header row
                set_cell_background(cell, "0F2942")
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="0F2942", sz="6")
                r = p_cell.add_run(text)
                r.font.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = WHITE
            else:
                # Data rows
                bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg_color)
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="E2E8F0", sz="4")
                r = p_cell.add_run(text)
                r.font.size = Pt(8.0)
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = NAVY
                else:
                    r.font.color.rgb = DARK_SLATE

    # ----------------------------------------------------
    # PRODUCTION LIFECYCLE & CALL TO ACTION
    # ----------------------------------------------------
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    cta_table = doc.add_table(rows=1, cols=1)
    cta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cta_table.autofit = False
    cta_cell = cta_table.cell(0, 0)
    cta_cell.width = Inches(7.0)
    set_cell_background(cta_cell, "0F2942")
    set_cell_margins(cta_cell, top=100, bottom=100, left=140, right=140)

    p_cta1 = cta_cell.paragraphs[0]
    p_cta1.paragraph_format.space_after = Pt(2)
    r_cta1 = p_cta1.add_run("Continuous Production Monitoring: Beyond Deployment")
    r_cta1.font.bold = True
    r_cta1.font.size = Pt(10)
    r_cta1.font.color.rgb = WHITE

    p_cta2 = cta_cell.add_paragraph()
    p_cta2.paragraph_format.space_after = Pt(4)
    p_cta2.paragraph_format.line_spacing = 1.15
    r_cta2 = p_cta2.add_run(
        "Our evaluations do not end at deployment. HashCase workflows operate with Shadow Deployments (auditing against human staff), "
        "Human-in-the-Loop confidence thresholds (routing borderline decisions to supervisors), and continuous drift monitors "
        "that alert our engineers the moment upstream LLM model updates or data distributions shift."
    )
    r_cta2.font.size = Pt(8.5)
    r_cta2.font.color.rgb = RGBColor(226, 232, 240)

    p_cta3 = cta_cell.add_paragraph()
    p_cta3.paragraph_format.space_after = Pt(0)
    r_cta3_bold = p_cta3.add_run("Deploy AI with Deterministic Confidence. ")
    r_cta3_bold.font.bold = True
    r_cta3_bold.font.size = Pt(8.5)
    r_cta3_bold.font.color.rgb = RGBColor(147, 197, 253) # Light Blue
    r_cta3_text = p_cta3.add_run(
        "Contact HashCase to schedule an AI Evaluation Audit or discuss hardening your automated workflows: "
    )
    r_cta3_text.font.size = Pt(8.5)
    r_cta3_text.font.color.rgb = WHITE
    r_contact = p_cta3.add_run("hello@hashcase.com  |  www.hashcase.com")
    r_contact.font.bold = True
    r_contact.font.size = Pt(8.5)
    r_contact.font.color.rgb = RGBColor(147, 197, 253)

    # Save document
    output_path = r"c:\Users\shaik\Desktop\dev-work\ai-newssharing\HashCase_Enterprise_AI_Evals_Capability_Brief.docx"
    doc.save(output_path)
    print(f"Document saved successfully at: {output_path}")

if __name__ == "__main__":
    create_document()
